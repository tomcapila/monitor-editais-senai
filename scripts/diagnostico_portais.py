"""
Diagnóstico de vários portais de uma vez, sem janela. Imprime um relatório
(tabelas, links com cara de edital, formulários, chamadas AJAX e, em portais
Paradigma, o teste do serviço JSON) e salva os arquivos em
diagnostico/portais/<site>/.

    python scripts/diagnostico_portais.py URL [URL ...] [--api URL ...]

--api consulta a URL direto (sem navegador) e mostra o que voltou: JSON,
planilha XLSX (com openpyxl instalado) ou texto. Serve para APIs de
transparência. É o que o workflow "Diagnóstico de portais" roda no GitHub Actions.
"""
import argparse
import io
import json
import re
import sys
from pathlib import Path

import httpx

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.coletores.base import USER_AGENT  # noqa: E402
from app.diagnostico import SAIDA, relatorio, rodar  # noqa: E402


def pasta(url: str) -> str:
    """Nome de pasta legível e único por URL (host + caminho + consulta)."""
    nome = re.sub(r"^https?://", "", url)
    return re.sub(r"[^A-Za-z0-9.-]+", "_", nome).strip("_")[:120]


def planilha(conteudo: bytes, max_linhas: int = 15) -> list[str]:
    """Primeiras linhas de cada aba de um XLSX."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ["(instale openpyxl para ler a planilha)"]
    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    linhas = []
    for ws in wb.worksheets:
        linhas.append(f"  aba {ws.title!r} ({ws.max_row} linhas x {ws.max_column} colunas)")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_linhas:
                break
            linhas.append(f"    {[str(v)[:80] if v is not None else '' for v in row]}")
    return linhas


PISTAS_JSON = ("credenciamento", "chamamento", "consultor", "instrutor", "instrutoria")


def resumo_lista(dados: list) -> list[str]:
    """Para uma lista de objetos JSON: campos, valores mais comuns dos campos
    com poucos valores distintos e os itens que parecem de consultoria."""
    from collections import Counter

    itens = [d for d in dados if isinstance(d, dict)]
    linhas = [f"  {len(dados)} itens; campos do primeiro: {sorted(itens[0]) if itens else []}"]
    for campo in sorted({k for d in itens for k in d}):
        valores = Counter(str(d.get(campo))[:60] for d in itens
                          if not isinstance(d.get(campo), (list, dict)))
        if 1 < len(valores) <= 40:
            linhas.append(f"  {campo}: {valores.most_common(15)}")
    achados = [d for d in itens
               if any(p in json.dumps(d, ensure_ascii=False).lower() for p in PISTAS_JSON)]
    linhas.append(f"  {len(achados)} itens citam {PISTAS_JSON}:")
    for d in achados[:40]:
        resumo = {k: (str(v)[:160] if not isinstance(v, (list, dict)) else f"[{len(v)}]")
                  for k, v in d.items() if v not in ("", None)}
        linhas.append(f"    {json.dumps(resumo, ensure_ascii=False)}")
    return linhas


def sondar(url: str, timeout: float = 60) -> str:
    """Consulta uma URL de API e descreve a resposta."""
    linhas = [f"=== API {url}"]
    try:
        r = httpx.get(url, timeout=timeout, follow_redirects=True,
                      headers={"User-Agent": USER_AGENT, "Accept": "application/json, */*"})
    except Exception as e:
        return "\n".join(linhas + [f"FALHOU: {type(e).__name__}: {e}"])
    tipo = r.headers.get("content-type", "")
    linhas.append(f"HTTP {r.status_code} [{tipo}] {len(r.content)} bytes; final: {r.url}")
    destino = SAIDA / "portais" / pasta(url)
    destino.mkdir(parents=True, exist_ok=True)
    (destino / "resposta.bin").write_bytes(r.content)
    if "json" in tipo:
        try:
            dados = r.json()
            if isinstance(dados, list):
                linhas += resumo_lista(dados)
            linhas.append(json.dumps(dados, ensure_ascii=False, indent=1)[:3000])
        except ValueError:
            linhas.append(r.text[:3000])
    elif r.content[:2] == b"PK":  # zip: XLSX (ODS não é lido)
        try:
            linhas += planilha(r.content)
        except Exception as e:
            linhas.append(f"(não abriu como XLSX: {type(e).__name__}: {e})")
    else:
        linhas.append(r.text[:3000])
    return "\n".join(linhas)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.strip().splitlines()[0])
    ap.add_argument("urls", nargs="*")
    ap.add_argument("--api", nargs="*", default=[], help="URLs consultadas sem navegador")
    ap.add_argument("--limite", type=float, default=180, help="segundos por portal")
    ap.add_argument("--todos-links", action="store_true",
                    help="lista todos os links, não só os com cara de edital")
    args = ap.parse_args()

    falhas = 0
    for url in args.urls:
        r = rodar(url, limite_segundos=args.limite, saida=SAIDA / "portais" / pasta(url),
                  amostra_max=6000)
        falhas += r["status"] != "ok"
        print(relatorio(r, todos_links=args.todos_links), end="\n\n", flush=True)
    for url in args.api:
        print(sondar(url), end="\n\n", flush=True)
    print(f"{len(args.urls) - falhas} de {len(args.urls)} portais diagnosticados; "
          f"{len(args.api)} APIs consultadas. Arquivos em {SAIDA / 'portais'}")
    return 1 if args.urls and falhas == len(args.urls) else 0


if __name__ == "__main__":
    sys.exit(main())
